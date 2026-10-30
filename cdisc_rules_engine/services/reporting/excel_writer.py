from io import BytesIO
from os import unlink
from tempfile import NamedTemporaryFile
from typing import List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.styles.cell_style import StyleArray
from openpyxl.worksheet.worksheet import Worksheet

_CELL_FONT = Font(name="Calibri", size=11)
_THIN_SIDE = Side(style="thin")
_CELL_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def excel_workbook_to_stream(wb):
    stream = None
    tmp = NamedTemporaryFile(delete=False)  # create tmp file, we'll controll when to delete it
    # close, NamedTemporaryFile() already opended it, we need it closed for wb.save()
    # or we'll get permission denied errors
    tmp.close()
    wb.save(tmp.name)  # save workbook to tmp file
    with open(tmp.name, "rb") as f:
        stream = f.read()  # open for reading
    unlink(tmp.name)  # delete
    return stream


def excel_update_worksheet(ws, rows, align_params=None, fill_empty_rows=False):
    for header_cell in ws[1]:
        if header_cell.value is not None:
            header_cell.border = _CELL_BORDER
    if not rows:
        return
    alignment = Alignment(**align_params) if align_params else None
    empty_fill = openpyxl.styles.PatternFill(start_color="ccffff", end_color="ccffff", fill_type="solid")
    styled = {}
    for row_num, row_data in enumerate(rows, 2):
        is_empty_row = fill_empty_rows and (row_data[1] == "" or row_data[1] is None)
        for col_num, col_data in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = col_data
            if is_empty_row:
                cell.fill = empty_fill
            _style_cell(cell, alignment, styled)


def _style_cell(cell, alignment: Optional[Alignment], styled: dict) -> None:
    key = tuple(cell._style) if cell.has_style else None
    cached = styled.get(key)
    if cached is None:
        cell.font = _CELL_FONT
        cell.border = _CELL_BORDER
        if alignment is not None:
            cell.alignment = alignment
        styled[key] = StyleArray(cell._style)
    else:
        cell._style = StyleArray(cached)


def excel_open_workbook(template_buffer):
    byte_stream = BytesIO(template_buffer)
    return openpyxl.load_workbook(byte_stream)


def group_rows_by_key(ws: Worksheet, rows: List[List], key_col: int = 0) -> None:
    if len(rows) < 2:
        return
    ws.sheet_properties.outlinePr.summaryBelow = False
    group_start = 0
    for i in range(1, len(rows) + 1):
        same_group = i < len(rows) and rows[i][key_col] == rows[group_start][key_col]
        if same_group:
            continue
        if i - group_start > 1:
            for row_num in range(group_start + 2 + 1, i + 2):
                ws.row_dimensions[row_num].outlineLevel = 1
        group_start = i
